#!/usr/bin/env python3
"""Generate the Bible Commerciale Session 1 PowerPoint presentation."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# --- Color palette (dark "Founder" style) ---
BG_DARK = RGBColor(0x0F, 0x0F, 0x14)
BG_CARD = RGBColor(0x1A, 0x1A, 0x24)
ACCENT = RGBColor(0x6C, 0x5C, 0xE7)       # Purple accent
ACCENT_LIGHT = RGBColor(0xA2, 0x96, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xB0, 0xB0, 0xC0)
MEDIUM_GRAY = RGBColor(0x80, 0x80, 0x96)
ORANGE = RGBColor(0xFD, 0x9F, 0x44)
GREEN = RGBColor(0x00, 0xD2, 0xA0)
RED_SOFT = RGBColor(0xFF, 0x6B, 0x6B)
YELLOW = RGBColor(0xFE, 0xCA, 0x57)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

W = prs.slide_width
H = prs.slide_height


def set_slide_bg(slide, color=BG_DARK):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, left, top, width, height, fill_color=None, border_color=None, border_width=Pt(0)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color or BG_CARD
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width
    else:
        shape.line.fill.background()
    # Rounded corners
    shape.adjustments[0] = 0.05
    return shape


def add_text(slide, left, top, width, height, text, font_size=18, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_list(slide, left, top, width, height, items, font_size=16, color=LIGHT_GRAY, bold_prefix=True):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(6)
        p.space_before = Pt(2)

        if isinstance(item, tuple):
            # (bold_part, normal_part)
            run1 = p.add_run()
            run1.text = item[0]
            run1.font.size = Pt(font_size)
            run1.font.color.rgb = WHITE
            run1.font.bold = True
            run1.font.name = "Calibri"
            run2 = p.add_run()
            run2.text = item[1]
            run2.font.size = Pt(font_size)
            run2.font.color.rgb = color
            run2.font.bold = False
            run2.font.name = "Calibri"
        else:
            run = p.add_run()
            run.text = f"  {item}"
            run.font.size = Pt(font_size)
            run.font.color.rgb = color
            run.font.bold = False
            run.font.name = "Calibri"
    return txBox


def add_quote(slide, left, top, width, height, text, font_size=18):
    # Accent bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.06), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    # Text
    txBox = slide.shapes.add_textbox(left + Inches(0.25), top, width - Inches(0.25), height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = ACCENT_LIGHT
    p.font.bold = False
    p.font.italic = True
    p.font.name = "Calibri"
    return txBox


def add_badge(slide, left, top, text, color=ACCENT):
    w, h = Inches(2.5), Inches(0.4)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.adjustments[0] = 0.3
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(12)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.CENTER
    return shape


# ============================================================
# SLIDE 1 — TITRE
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_slide_bg(slide)

# Central accent line
line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.5), Inches(2.2), Inches(2.3), Inches(0.06))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT
line.line.fill.background()

add_text(slide, Inches(1), Inches(2.5), Inches(11.3), Inches(1.2),
         "BIBLE COMMERCIALE", font_size=48, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(3.6), Inches(11.3), Inches(0.8),
         "Session 1 — Les Fondations de la Conviction", font_size=28, color=ACCENT_LIGHT, bold=False, alignment=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(5.0), Inches(11.3), Inches(0.5),
         "Incubateur UPVD", font_size=20, color=MEDIUM_GRAY, bold=False, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 2 — PROMESSE DE LA MATINÉE
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_badge(slide, Inches(0.8), Inches(0.6), "PROMESSE")
add_text(slide, Inches(0.8), Inches(1.2), Inches(11.5), Inches(1),
         "Ce matin, on ne produit pas un pitch générique.", font_size=32, color=WHITE, bold=True)
add_text(slide, Inches(0.8), Inches(2.1), Inches(11.5), Inches(0.8),
         "On construit la première version de votre Bible Commerciale.", font_size=26, color=ACCENT_LIGHT, bold=False)

add_text(slide, Inches(0.8), Inches(3.3), Inches(11.5), Inches(0.5),
         "À la fin de la session, vous aurez clarifié :", font_size=20, color=LIGHT_GRAY)

items = [
    "Qui vous devez convaincre",
    "Pourquoi ces personnes disent oui",
    "Quels sont leurs vrais enjeux",
    "Comment relier votre startup à ces enjeux",
]
for i, item in enumerate(items):
    y = Inches(4.1) + Inches(i * 0.55)
    # Accent dot
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), y + Inches(0.08), Inches(0.15), Inches(0.15))
    dot.fill.solid()
    dot.fill.fore_color.rgb = ACCENT
    dot.line.fill.background()
    add_text(slide, Inches(1.6), y, Inches(10), Inches(0.5), item, font_size=20, color=WHITE)

# ============================================================
# SLIDE 3 — ACCUEIL / CONNEXION
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_badge(slide, Inches(0.8), Inches(0.6), "9h00 — 9h15")
add_text(slide, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.8),
         "Accueil, connexion, mise en route", font_size=36, color=WHITE, bold=True)

# Card 1
add_shape(slide, Inches(0.8), Inches(2.5), Inches(5.5), Inches(2.2))
add_text(slide, Inches(1.1), Inches(2.7), Inches(5), Inches(0.4),
         "INSCRIPTION", font_size=14, color=ACCENT, bold=True)
add_bullet_list(slide, Inches(1.1), Inches(3.2), Inches(5), Inches(1.5), [
    "Nom de la startup",
    "Nom du fondateur",
    "Mot de passe",
    "Secteur & stade",
], font_size=16)

# Card 2
add_shape(slide, Inches(7), Inches(2.5), Inches(5.5), Inches(2.2))
add_text(slide, Inches(7.3), Inches(2.7), Inches(5), Inches(0.4),
         "TOUR DE TABLE", font_size=14, color=ACCENT, bold=True)
add_bullet_list(slide, Inches(7.3), Inches(3.2), Inches(5), Inches(1.5), [
    "Votre nom",
    "Votre startup",
    "Ce que vous devez obtenir dans les 90 jours",
], font_size=16)

add_quote(slide, Inches(0.8), Inches(5.3), Inches(11.5), Inches(1),
          "« Tout ce qu'on construit sera saisi dans l'application pour produire votre Bible Commerciale v1. »")

# ============================================================
# SLIDE 4 — CONVAINCRE C'EST VOTRE MÉTIER
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_badge(slide, Inches(0.8), Inches(0.6), "9h15 — 9h35")
add_text(slide, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.8),
         "Convaincre, c'est déjà votre métier", font_size=36, color=WHITE, bold=True)

# 5 cards
roles = [
    ("Vision", "Accompagnateurs", ACCENT),
    ("Projet", "Talents / Équipe", GREEN),
    ("Solution", "Clients", ORANGE),
    ("Potentiel", "Financeurs", YELLOW),
    ("Crédibilité", "Partenaires", RED_SOFT),
]
card_w = Inches(2.2)
start_x = Inches(0.8)
gap = Inches(0.25)
for i, (what, who, col) in enumerate(roles):
    x = start_x + i * (card_w + gap)
    shape = add_shape(slide, x, Inches(2.5), card_w, Inches(2))
    # Color top bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(2.5), card_w, Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = col
    bar.line.fill.background()
    add_text(slide, x + Inches(0.2), Inches(2.8), card_w - Inches(0.4), Inches(0.5),
             what, font_size=22, color=col, bold=True, alignment=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.2), Inches(3.4), card_w - Inches(0.4), Inches(0.8),
             f"→ {who}", font_size=16, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

add_quote(slide, Inches(0.8), Inches(5.2), Inches(11.5), Inches(1.2),
          "« Le sujet n'est pas comment vendre mon offre. Le sujet est : comment faire avancer des décisions favorables dans mon écosystème ? »")

# Message clé
add_shape(slide, Inches(3), Inches(6.5), Inches(7.3), Inches(0.6), fill_color=ACCENT)
add_text(slide, Inches(3), Inches(6.5), Inches(7.3), Inches(0.6),
         "Vous vendez déjà. La question : le faites-vous consciemment et méthodiquement ?",
         font_size=16, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 5 — MODULE 1 : CARTOGRAPHIE
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_badge(slide, Inches(0.8), Inches(0.6), "MODULE 1", color=GREEN)
add_text(slide, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.8),
         "Cartographie des parties prenantes", font_size=36, color=WHITE, bold=True)

# 6 category cards (2 rows of 3)
categories = [
    ("Accompagnateurs", "Incubateur, mentor, référent"),
    ("Équipe", "Associé, 1er salarié, freelance clé"),
    ("Clients", "Bêta-testeur, 1er client payant"),
    ("Financeurs", "BA, BPI, banque, subvention"),
    ("Partenaires", "Distributeur, apporteur d'affaires"),
    ("Écosystème", "Média local, réseau, asso, event"),
]
card_w = Inches(3.7)
card_h = Inches(1.5)
gap_x = Inches(0.3)
gap_y = Inches(0.2)
for i, (cat, ex) in enumerate(categories):
    row = i // 3
    col = i % 3
    x = Inches(0.8) + col * (card_w + gap_x)
    y = Inches(2.5) + row * (card_h + gap_y)
    add_shape(slide, x, y, card_w, card_h)
    add_text(slide, x + Inches(0.3), y + Inches(0.2), card_w - Inches(0.6), Inches(0.4),
             cat, font_size=18, color=ACCENT_LIGHT, bold=True)
    add_text(slide, x + Inches(0.3), y + Inches(0.7), card_w - Inches(0.6), Inches(0.6),
             ex, font_size=14, color=MEDIUM_GRAY)

# Bottom rule
add_shape(slide, Inches(0.8), Inches(6.1), Inches(11.5), Inches(0.9), fill_color=RGBColor(0x15, 0x15, 0x20))
add_text(slide, Inches(1.1), Inches(6.2), Inches(11), Inches(0.7),
         "Chaque entrée : Nom + Rôle + Priorité (critique / important / secondaire)  —  Minimum 3 catégories remplies",
         font_size=16, color=ORANGE, bold=True, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 6 — SORTEZ DU MOT MARCHÉ
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_text(slide, Inches(1.5), Inches(1.5), Inches(10.3), Inches(1.5),
         "« Votre première erreur, c'est de croire que votre marché suffit. »",
         font_size=36, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.5), Inches(3.2), Inches(2.3), Inches(0.04))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT
line.line.fill.background()

add_text(slide, Inches(1.5), Inches(3.6), Inches(10.3), Inches(1.2),
         "Ce qui compte, c'est votre système concret\nde parties prenantes.",
         font_size=28, color=ACCENT_LIGHT, alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(1.5), Inches(5.2), Inches(10.3), Inches(1),
         "Ne restez pas enfermés dans « mes clients ».\nVoyez l'ensemble des personnes qui peuvent accélérer ou freiner votre startup.",
         font_size=20, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 7 — DEBRIEF PRIORITÉS
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_badge(slide, Inches(0.8), Inches(0.6), "DEBRIEF", color=ORANGE)
add_text(slide, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.8),
         "Où mettez-vous votre énergie ?", font_size=36, color=WHITE, bold=True)

# 3 priority cards
priorities = [
    ("CRITIQUE", "Énergie maximale", RED_SOFT, "Les parties prenantes\nsans lesquelles rien\nn'avance"),
    ("IMPORTANT", "Énergie ciblée", ORANGE, "Celles qui accélèrent\nsignificativement\nvotre trajectoire"),
    ("SECONDAIRE", "Énergie mesurée", MEDIUM_GRAY, "Utiles mais pas\ndécisives à court\nterme"),
]
for i, (label, sub, col, desc) in enumerate(priorities):
    x = Inches(0.8) + i * Inches(4.1)
    add_shape(slide, x, Inches(2.5), Inches(3.8), Inches(2.5), border_color=col, border_width=Pt(2))
    add_text(slide, x + Inches(0.3), Inches(2.7), Inches(3.2), Inches(0.5),
             label, font_size=22, color=col, bold=True, alignment=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.3), Inches(3.2), Inches(3.2), Inches(0.4),
             sub, font_size=14, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.3), Inches(3.7), Inches(3.2), Inches(1),
             desc, font_size=14, color=MEDIUM_GRAY, alignment=PP_ALIGN.CENTER)

add_quote(slide, Inches(0.8), Inches(5.6), Inches(11.5), Inches(1),
          "« Qui avez-vous mis en critique, et pourquoi ? »", font_size=22)

# ============================================================
# SLIDE 8 — KAHNEMAN
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_badge(slide, Inches(0.8), Inches(0.6), "MODULE 2", color=ACCENT)
add_text(slide, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.8),
         "Pourquoi les gens disent oui — Kahneman", font_size=36, color=WHITE, bold=True)

# System 1
add_shape(slide, Inches(0.8), Inches(2.5), Inches(5.5), Inches(2.8), border_color=ORANGE, border_width=Pt(2))
add_text(slide, Inches(1.1), Inches(2.7), Inches(5), Inches(0.5),
         "SYSTÈME 1", font_size=26, color=ORANGE, bold=True, alignment=PP_ALIGN.CENTER)
add_text(slide, Inches(1.1), Inches(3.3), Inches(5), Inches(0.5),
         "Rapide • Intuitif • Émotionnel", font_size=18, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
add_bullet_list(slide, Inches(1.5), Inches(3.9), Inches(4.5), Inches(1.2), [
    "Confiance",
    "Envie",
    "Alignement",
], font_size=16, color=ORANGE)

# System 2
add_shape(slide, Inches(7), Inches(2.5), Inches(5.5), Inches(2.8), border_color=GREEN, border_width=Pt(2))
add_text(slide, Inches(7.3), Inches(2.7), Inches(5), Inches(0.5),
         "SYSTÈME 2", font_size=26, color=GREEN, bold=True, alignment=PP_ALIGN.CENTER)
add_text(slide, Inches(7.3), Inches(3.3), Inches(5), Inches(0.5),
         "Lent • Analytique • Logique", font_size=18, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
add_bullet_list(slide, Inches(7.7), Inches(3.9), Inches(4.5), Inches(1.2), [
    "Preuve",
    "Cohérence",
    "Logique",
], font_size=16, color=GREEN)

add_quote(slide, Inches(0.8), Inches(5.8), Inches(11.5), Inches(1),
          "« Les gens ressentent d'abord, puis rationalisent ensuite. Votre stratégie doit adresser les deux. »")

# ============================================================
# SLIDE 9 — MINI QUIZ
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_badge(slide, Inches(0.8), Inches(0.6), "QUIZ", color=YELLOW)
add_text(slide, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.8),
         "Mini-quiz : Système 1 ou Système 2 ?", font_size=36, color=WHITE, bold=True)

add_text(slide, Inches(1.5), Inches(2.8), Inches(10.3), Inches(1.5),
         "Le quiz n'est pas un contrôle scolaire.\n\nIl sert à prendre conscience des mécanismes mentaux\nque vous activez — ou oubliez d'activer —\nquand vous présentez votre startup.",
         font_size=22, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# Bilan box
add_shape(slide, Inches(2.5), Inches(5.0), Inches(8.3), Inches(1.8))
add_text(slide, Inches(2.8), Inches(5.1), Inches(7.7), Inches(0.4),
         "BILAN INTERMÉDIAIRE", font_size=14, color=ACCENT, bold=True)
add_bullet_list(slide, Inches(2.8), Inches(5.5), Inches(7.7), Inches(1.2), [
    ("✓ ", "Vous avez identifié QUI convaincre"),
    ("✓ ", "Vous comprenez COMMENT ces personnes décident"),
    ("→ ", "Prochaine étape : leurs VRAIS ENJEUX"),
], font_size=16)

# ============================================================
# SLIDE 10 — PAUSE
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_text(slide, Inches(1), Inches(2.5), Inches(11.3), Inches(1),
         "PAUSE", font_size=60, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(3.8), Inches(11.3), Inches(0.6),
         "11h00 — 11h15", font_size=28, color=MEDIUM_GRAY, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 11 — LES 8 BIAIS COGNITIFS
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_badge(slide, Inches(0.8), Inches(0.6), "MODULE 2 (suite)", color=ACCENT)
add_text(slide, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.8),
         "Les 8 biais cognitifs", font_size=36, color=WHITE, bold=True)

biais = [
    ("Ancrage", "La 1ère info conditionne le jugement"),
    ("Confirmation", "On cherche ce qui confirme nos croyances"),
    ("Preuve sociale", "On suit ce que font les autres"),
    ("Aversion à la perte", "Perdre fait plus mal que gagner"),
    ("Statu quo", "On préfère ne pas changer"),
    ("Halo", "Une impression positive irradie sur tout"),
    ("Rareté", "Ce qui est rare est perçu comme précieux"),
    ("Réciprocité", "On rend ce qu'on reçoit"),
]
colors_biais = [ACCENT, GREEN, ORANGE, RED_SOFT, MEDIUM_GRAY, YELLOW, ACCENT_LIGHT, GREEN]
card_w = Inches(2.8)
card_h = Inches(1.8)
gap_x = Inches(0.2)
gap_y = Inches(0.2)
for i, ((name, desc), col) in enumerate(zip(biais, colors_biais)):
    row = i // 4
    c = i % 4
    x = Inches(0.5) + c * (card_w + gap_x)
    y = Inches(2.3) + row * (card_h + gap_y)
    add_shape(slide, x, y, card_w, card_h, border_color=col, border_width=Pt(1.5))
    # Number
    add_text(slide, x + Inches(0.15), y + Inches(0.1), Inches(0.4), Inches(0.4),
             str(i + 1), font_size=14, color=col, bold=True)
    add_text(slide, x + Inches(0.15), y + Inches(0.45), card_w - Inches(0.3), Inches(0.45),
             name, font_size=17, color=WHITE, bold=True)
    add_text(slide, x + Inches(0.15), y + Inches(0.95), card_w - Inches(0.3), Inches(0.7),
             desc, font_size=12, color=MEDIUM_GRAY)

# ============================================================
# SLIDE 12 — EXERCICE BIAIS × PP
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_badge(slide, Inches(0.8), Inches(0.6), "EXERCICE", color=ORANGE)
add_text(slide, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.8),
         "Appliquer les biais à vos parties prenantes critiques", font_size=32, color=WHITE, bold=True)

add_text(slide, Inches(0.8), Inches(2.3), Inches(11.5), Inches(0.6),
         "Pour vos 3 PP prioritaires → sélectionnez 3 biais max → écrivez comment les activer concrètement.",
         font_size=20, color=LIGHT_GRAY)

# Questions de relance
add_shape(slide, Inches(0.8), Inches(3.2), Inches(5.5), Inches(3))
add_text(slide, Inches(1.1), Inches(3.3), Inches(5), Inches(0.4),
         "QUESTIONS DE RELANCE", font_size=14, color=ACCENT, bold=True)
add_bullet_list(slide, Inches(1.1), Inches(3.8), Inches(5), Inches(2.2), [
    "Quelle preuve sociale pouvez-vous montrer ?",
    "Quel coût de l'inaction pouvez-vous révéler ?",
    "Quel effort de changement devez-vous réduire ?",
    "Quel 1er signal de crédibilité devez-vous soigner ?",
], font_size=16)

# Cadrage éthique
add_shape(slide, Inches(7), Inches(3.2), Inches(5.5), Inches(3), border_color=ACCENT, border_width=Pt(1.5))
add_text(slide, Inches(7.3), Inches(3.3), Inches(5), Inches(0.4),
         "CADRAGE ÉTHIQUE", font_size=14, color=ACCENT, bold=True)
add_text(slide, Inches(7.3), Inches(3.9), Inches(5), Inches(2),
         "On ne manipule pas.\n\nOn comprend les mécanismes de décision pour mieux construire une relation de conviction.\n\nUn financeur, un accompagnateur et un premier client ne disent pas oui pour les mêmes raisons.",
         font_size=16, color=LIGHT_GRAY)

# ============================================================
# SLIDE 13 — TRANSITION ENJEUX
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_text(slide, Inches(1), Inches(1.5), Inches(11.3), Inches(0.7),
         "Jusqu'ici, on a répondu à deux questions :", font_size=24, color=MEDIUM_GRAY, alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(1), Inches(2.5), Inches(11.3), Inches(0.6),
         "✓  Qui faut-il convaincre ?", font_size=28, color=GREEN, bold=True, alignment=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(3.2), Inches(11.3), Inches(0.6),
         "✓  Comment ces personnes disent-elles oui ?", font_size=28, color=GREEN, bold=True, alignment=PP_ALIGN.CENTER)

line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.5), Inches(4.2), Inches(2.3), Inches(0.04))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT
line.line.fill.background()

add_text(slide, Inches(1), Inches(4.6), Inches(11.3), Inches(0.8),
         "Maintenant, la question décisive :", font_size=24, color=MEDIUM_GRAY, alignment=PP_ALIGN.CENTER)
add_text(slide, Inches(1), Inches(5.4), Inches(11.3), Inches(0.8),
         "Qu'est-ce qui compte vraiment pour elles ?", font_size=36, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 14 — MODULE 3 : MATRICE DES ENJEUX
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_badge(slide, Inches(0.8), Inches(0.6), "MODULE 3", color=GREEN)
add_text(slide, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.8),
         "La Matrice des Enjeux", font_size=36, color=WHITE, bold=True)

add_text(slide, Inches(0.8), Inches(2.1), Inches(3), Inches(0.5),
         "C'est le cœur de la session.", font_size=18, color=ACCENT_LIGHT, bold=True)

# 3 columns
cols_data = [
    ("ENJEU APPARENT", "Ce que la personne semble vouloir dans la relation avec la startup", ORANGE),
    ("ENJEU PROFOND", "Ce qui l'anime vraiment dans son propre contexte", RED_SOFT),
    ("PONT", "Le lien logique ET émotionnel entre votre startup et cet enjeu", GREEN),
]
for i, (title, desc, col) in enumerate(cols_data):
    x = Inches(0.8) + i * Inches(4.1)
    add_shape(slide, x, Inches(2.8), Inches(3.8), Inches(2.2), border_color=col, border_width=Pt(2))
    # Top colored bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(2.8), Inches(3.8), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = col
    bar.line.fill.background()
    add_text(slide, x + Inches(0.3), Inches(3.0), Inches(3.2), Inches(0.5),
             title, font_size=18, color=col, bold=True, alignment=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.3), Inches(3.6), Inches(3.2), Inches(1.2),
             desc, font_size=15, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# Questions
add_shape(slide, Inches(0.8), Inches(5.4), Inches(11.5), Inches(1.4))
add_text(slide, Inches(1.1), Inches(5.5), Inches(5), Inches(0.3),
         "QUESTIONS CLÉ", font_size=14, color=ACCENT, bold=True)
add_bullet_list(slide, Inches(1.1), Inches(5.9), Inches(10.5), Inches(0.8), [
    "À qui rend-elle des comptes ?       •   Qu'est-ce qui l'empêche de dormir ?       •   Quels sont ses KPIs ?",
], font_size=15)

# ============================================================
# SLIDE 15 — EXEMPLE ACCOMPAGNATEUR
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_badge(slide, Inches(0.8), Inches(0.6), "EXEMPLE", color=YELLOW)
add_text(slide, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.8),
         "L'accompagnateur d'incubateur", font_size=36, color=WHITE, bold=True)

# 3 cards
examples = [
    ("ENJEU APPARENT", "Vous aider à réussir", ORANGE),
    ("ENJEU PROFOND", "Démontrer la réussite de sa cohorte\nJustifier les financements du programme\nCréer de l'emploi local\nRenforcer sa propre crédibilité", RED_SOFT),
    ("PONT", "Votre traction rapide est\nSA meilleure preuve de valeur", GREEN),
]
for i, (title, content, col) in enumerate(examples):
    x = Inches(0.8) + i * Inches(4.1)
    h = Inches(3) if i == 1 else Inches(2.2)
    add_shape(slide, x, Inches(2.5), Inches(3.8), Inches(3), border_color=col, border_width=Pt(2))
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(2.5), Inches(3.8), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = col
    bar.line.fill.background()
    add_text(slide, x + Inches(0.3), Inches(2.7), Inches(3.2), Inches(0.4),
             title, font_size=14, color=col, bold=True)
    add_text(slide, x + Inches(0.3), Inches(3.2), Inches(3.2), Inches(2),
             content, font_size=16, color=LIGHT_GRAY)

add_quote(slide, Inches(0.8), Inches(6.0), Inches(11.5), Inches(0.8),
          "« Ne remplissez pas ça comme un exercice littéraire. Remplissez-le comme un outil de lecture du réel. »")

# ============================================================
# SLIDE 16 — MODULE 4 : BIBLE V1
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_badge(slide, Inches(0.8), Inches(0.6), "MODULE 4", color=ACCENT)
add_text(slide, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.8),
         "Votre Bible Commerciale v1", font_size=36, color=WHITE, bold=True)

# 3 synthesis cards
synth = [
    ("Carte des\nParties Prenantes", ACCENT),
    ("Matrice\ndes Enjeux", GREEN),
    ("Leviers\nCognitifs", ORANGE),
]
for i, (label, col) in enumerate(synth):
    x = Inches(0.8) + i * Inches(4.1)
    shape = add_shape(slide, x, Inches(2.5), Inches(3.8), Inches(1.8), border_color=col, border_width=Pt(2))
    add_text(slide, x + Inches(0.3), Inches(2.8), Inches(3.2), Inches(1.2),
             label, font_size=24, color=col, bold=True, alignment=PP_ALIGN.CENTER)

# Arrow + PDF
add_text(slide, Inches(1), Inches(4.7), Inches(11.3), Inches(0.5),
         "↓", font_size=30, color=ACCENT, alignment=PP_ALIGN.CENTER)

add_shape(slide, Inches(3.5), Inches(5.2), Inches(6.3), Inches(0.7), fill_color=ACCENT)
add_text(slide, Inches(3.5), Inches(5.2), Inches(6.3), Inches(0.7),
         "Export PDF  →  bible-commerciale-[startup]-session1.pdf",
         font_size=18, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

add_quote(slide, Inches(1.5), Inches(6.3), Inches(10.3), Inches(0.8),
          "« Vous ne repartez pas avec des conseils. Vous repartez avec une première version de votre Bible Commerciale. »")

# ============================================================
# SLIDE 17 — SUITE DU PROGRAMME
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_text(slide, Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.8),
         "La suite du programme", font_size=36, color=WHITE, bold=True)

sessions = [
    ("S2", "Les Ponts d'Enjeux", ACCENT),
    ("S3", "Les Canaux d'Accès", GREEN),
    ("S4", "Structurer l'Entretien", ORANGE),
    ("S5-6", "Découverte & Maïeutique", YELLOW),
    ("S7", "Argumentation, Objections & Closing", RED_SOFT),
    ("S8-9", "DISC", ACCENT_LIGHT),
    ("S10", "Synthèse complète", WHITE),
]
for i, (num, title, col) in enumerate(sessions):
    y = Inches(1.9) + i * Inches(0.7)
    # Number badge
    badge_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), y, Inches(1), Inches(0.45))
    badge_shape.fill.solid()
    badge_shape.fill.fore_color.rgb = col
    badge_shape.line.fill.background()
    badge_shape.adjustments[0] = 0.3
    tf = badge_shape.text_frame
    p = tf.paragraphs[0]
    p.text = num
    p.font.size = Pt(14)
    p.font.color.rgb = BG_DARK if col in [WHITE, YELLOW] else WHITE
    p.font.bold = True
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.CENTER
    # Title
    add_text(slide, Inches(2.8), y, Inches(8), Inches(0.45),
             title, font_size=20, color=LIGHT_GRAY)

# ============================================================
# SLIDE 18 — ENGAGEMENTS
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)

add_text(slide, Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.8),
         "Vos engagements avant la prochaine session", font_size=36, color=WHITE, bold=True)

engagements = [
    ("3", "parties prenantes à travailler en priorité", ACCENT),
    ("1", "hypothèse d'enjeu à vérifier", GREEN),
    ("1", "levier cognitif à tester", ORANGE),
    ("1", "conversation à provoquer", RED_SOFT),
]
for i, (num, text, col) in enumerate(engagements):
    y = Inches(2.2) + i * Inches(1.2)
    # Big number
    add_text(slide, Inches(1.5), y, Inches(1), Inches(1),
             num, font_size=48, color=col, bold=True, alignment=PP_ALIGN.CENTER)
    # Text
    add_text(slide, Inches(2.8), y + Inches(0.15), Inches(8), Inches(0.7),
             text, font_size=24, color=LIGHT_GRAY)

add_quote(slide, Inches(1.5), Inches(6.2), Inches(10.3), Inches(0.8),
          "« Aujourd'hui, vous avez posé les fondations de la conviction. »", font_size=22)

# ============================================================
# SAVE
# ============================================================
output_path = "/home/user/upvd-2026/public/bible-commerciale-session1-support.pptx"
prs.save(output_path)
print(f"Presentation saved to {output_path}")
print(f"Total slides: {len(prs.slides)}")
